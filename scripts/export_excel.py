# -*- coding: utf-8 -*-
"""
Exporta TODAS las tablas de la BD (BD/mrpeanutt_costos.db) a un solo Excel de consulta,
una hoja por tabla: CENTRO_DE_COSTOS_MRPEANUTT.xlsx (raíz del proyecto).

Correr al final de la cadena:

    python etl_build_db.py
    python export_dashboard_json.py
    python build_dashboard.py
    python export_excel.py
"""
import json
import re
import sqlite3
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
DB_PATH = BASE / "BD" / "mrpeanutt_costos.db"
DATA_JSON = BASE / "dashboard" / "data.json"
OUT = BASE / "CENTRO_DE_COSTOS_MRPEANUTT.xlsx"

# (hoja, grupo, tabla en la BD, descripción, ORDER BY)
HOJAS = [
    ("02 Parametros", "Base", "parametros", "Insumos del Excel y parámetros de la propuesta (IGV, costos por kg, PVP almendra, porción)", "id"),
    ("03 Productos", "Base", "productos_costeo", "Costeo de las 8 presentaciones: 4 frascos 150 g, baldes de maní 4 kg y 1 kg, baldes de almendra 4 kg y 1 kg (simulación)", "id"),
    ("04 Costeo baldes", "Base", "costeo_balde_detalle", "Desglose del costo de cada balde: envase (compra real), mantequilla y etiqueta", "id"),
    ("04b Compra envases", "Base", "compras_envases", "Compra real de envases de balde (12 x S/36 y 12 x S/55) y diferencia con el Excel", "id"),
    ("05 Canal directo", "Directo", "canal_directo", "PVP al público (sin IGV, NRUS), ganancia y margen de cada frasco", "id"),
    ("06 Promociones Excel", "Directo", "promociones", "Promos del Excel (2 y 3 almendras, combo 3 mantequillas): precio, descuento real y margen", "id"),
    ("06b Packs propuesta", "Directo", "promociones_propuesta", "PROPUESTA: mecánica 'arma tu pack' (precio fijo por cantidad) con rango de descuento y margen sobre todas las mezclas", "id"),
    ("06c Packs mezclas", "Directo", "promociones_ejemplos", "Todas las mezclas posibles de cada pack con su descuento, ahorro y margen", "id"),
    ("06d Packs terminos", "Directo", "promociones_terminos", "Términos y condiciones sugeridos de la mecánica de packs y por qué", "orden"),
    ("06f Packs VIP", "Directo", "packs_vip", "Packs VIP (solo clientes fidelizados): 2x S/35, 3x S/45, 4x S/60, 2 almendras S/42; rango de margen sobre todas las mezclas", "id"),
    ("06g Packs VIP mezclas", "Directo", "packs_vip_mezclas", "Todas las mezclas posibles de cada pack VIP con su margen", "id"),
    ("06e Packs escenarios", "Directo", "packs_escenarios", "Precio de cada pack según el margen objetivo (65/60/55/50/45%) sobre la peor mezcla", "id"),
    ("07 Escala B2B Excel", "B2B", "escala_b2b_excel", "Escala por tramo tal como está en el Excel, con el margen que se lleva la tienda al PVP", "id"),
    ("08 Escala B2B propuesta", "B2B", "escala_b2b_propuesta", "PROPUESTA: tramos sin huecos, precios redondeados, margen Mr. Peanutt y margen de la tienda", "id"),
    ("09 Cadena de valor", "B2B", "cadena_valor", "Quién gana qué (Mr. Peanutt / distribuidor / tienda) en cada escenario de venta", "id"),
    ("10 Restaurantes porcion", "B2B", "restaurantes_porcion", "Costo por kg y por porción de 30 g de cada formato y tramo, ahorro vs frasco", "id"),
    ("10b Baldes publico", "Directo", "baldes_publico", "Baldes a público final (sin IGV): precio recomendado, margen y comparación con restaurante y frasco", "id"),
    ("10c Baldes publico escenarios", "Directo", "baldes_publico_escenarios", "Precio de cada balde a público final según el margen (50-70%)", "id"),
    ("11 Resumen margenes", "Resumen", "resumen_margenes", "Rango de márgenes de Mr. Peanutt y de la tienda por canal y producto", "id"),
    ("11b Costos fijos", "Equilibrio", "costos_fijos", "Costos fijos mensuales: sueldo S/1,020 + NRUS S/50 + Cloud USD 20", "id"),
    ("11c Punto equilibrio", "Equilibrio", "punto_equilibrio", "Venta mensual que cubre los costos fijos, por escenario de canal y el conservador (margen más bajo)", "id"),
    ("12 Recomendaciones", "Datos", "recomendaciones", "Decisiones sugeridas, en orden de prioridad", "prioridad"),
    ("13 Historial", "Historial", "issues", "Hallazgos del Excel: severidad, estado y cómo se resolvieron",
     "CASE severidad WHEN 'alta' THEN 0 WHEN 'media' THEN 1 ELSE 2 END, id"),
]

CACAO = "4A2C17"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=CACAO)
TITLE_FONT = Font(size=12, color=CACAO)
FMT_SOLES = '"S/ "#,##0.00'
FMT_PCT = "0.0%"

RE_PCT = re.compile(r"(pct|_share$|^share|porcentaje)", re.I)
RE_SOLES = re.compile(
    r"(monto|venta|costo|precio|pago|ganancia|gana_|contribucion|igv_soles|compra_minima|pvp|suma_)", re.I)
RE_NO_SOLES = re.compile(r"(unidades|numero|_n$|^id$|tramo|cantidad|peso|porcion_gr|prioridad|valor)", re.I)


def formato_columna(nombre, valores):
    if not any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in valores):
        return None
    if RE_PCT.search(nombre):
        return FMT_PCT
    if RE_NO_SOLES.search(nombre):
        return "#,##0.##"
    if RE_SOLES.search(nombre):
        return FMT_SOLES
    return None


def escribir_tabla(ws, titulo, columnas, filas):
    ws["A1"] = titulo
    ws["A1"].font = TITLE_FONT
    for j, col in enumerate(columnas, start=1):
        c = ws.cell(row=2, column=j, value=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    for i, fila in enumerate(filas, start=3):
        for j, v in enumerate(fila, start=1):
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "A3"
    for j, col in enumerate(columnas, start=1):
        valores = [f[j - 1] for f in filas]
        fmt = formato_columna(col, valores)
        if fmt:
            for i in range(3, 3 + len(filas)):
                ws.cell(row=i, column=j).number_format = fmt
        largo = max([len(str(col))] + [len(str(v)) for v in valores if v is not None] or [10])
        ws.column_dimensions[get_column_letter(j)].width = min(max(largo + 2, 10), 60)


def main():
    cx = sqlite3.connect(DB_PATH)
    kpis = json.loads(DATA_JSON.read_text(encoding="utf-8"))["kpis"]
    wb = openpyxl.Workbook()

    ws_idx = wb.active
    ws_idx.title = "00 Índice"
    ws_kpi = wb.create_sheet("01 KPIs")

    indice = [("01 KPIs", "Resumen", "Indicadores principales del centro de costos", 1)]
    for hoja, grupo, tabla, desc, orden in HOJAS:
        cur = cx.execute(f"SELECT * FROM {tabla} ORDER BY {orden}")
        columnas = [d[0] for d in cur.description]
        filas = [list(r) for r in cur.fetchall()]
        ws = wb.create_sheet(hoja)
        escribir_tabla(ws, desc, columnas, filas)
        indice.append((hoja, grupo, desc, len(filas)))

    # -- 01 KPIs --
    ws_kpi["A1"] = "INDICADORES PRINCIPALES"
    ws_kpi["A1"].font = TITLE_FONT
    kpi_rows = [
        ("Presentaciones costeadas", kpis["n_productos"], "0"),
        ("Costos fijos mensuales", kpis["costos_fijos_total"], FMT_SOLES),
        ("Punto de equilibrio conservador (sin IGV)", kpis["pe_conservador_sin_igv"], FMT_SOLES),
        ("Punto de equilibrio conservador (con IGV)", kpis["pe_conservador_con_igv"], FMT_SOLES),
        ("   margen usado (el más bajo de la escala)", kpis["pe_conservador_margen"], FMT_PCT),
        ("Punto de equilibrio solo canal directo, maní (con IGV)", kpis["pe_directo_mani_con_igv"], FMT_SOLES),
        ("   frascos de maní al mes", kpis["pe_directo_mani_unidades"], "#,##0"),
        ("Margen canal directo (sin IGV, NRUS) - mínimo", kpis["margen_directo_min"], FMT_PCT),
        ("Margen canal directo (sin IGV, NRUS) - máximo", kpis["margen_directo_max"], FMT_PCT),
        ("Promociones propuestas: margen mínimo", kpis["promos_margen_min"], FMT_PCT),
        ("Promociones propuestas: margen máximo", kpis["promos_margen_max"], FMT_PCT),
        ("Margen Mr. Peanutt tiendas (Excel) - mínimo", kpis["margen_tiendas_excel_min"], FMT_PCT),
        ("Margen Mr. Peanutt tiendas (Excel) - máximo", kpis["margen_tiendas_excel_max"], FMT_PCT),
        ("Margen de la tienda al PVP (Excel) - mínimo", kpis["margen_tienda_excel_min"], FMT_PCT),
        ("Margen de la tienda al PVP (Excel) - máximo", kpis["margen_tienda_excel_max"], FMT_PCT),
        ("Margen Mr. Peanutt tiendas (propuesta) - mínimo", kpis["margen_tiendas_prop_min"], FMT_PCT),
        ("Margen Mr. Peanutt tiendas (propuesta) - máximo", kpis["margen_tiendas_prop_max"], FMT_PCT),
        ("Margen de la tienda al PVP (propuesta) - mínimo", kpis["margen_tienda_prop_min"], FMT_PCT),
        ("Margen de la tienda al PVP (propuesta) - máximo", kpis["margen_tienda_prop_max"], FMT_PCT),
        ("Margen Mr. Peanutt restaurantes (balde 4 kg) - mínimo", kpis["margen_restaurantes_min"], FMT_PCT),
        ("Margen Mr. Peanutt restaurantes (balde 4 kg) - máximo", kpis["margen_restaurantes_max"], FMT_PCT),
        ("Balde 4 kg: precio por kg con IGV - mínimo", kpis["balde4_precio_kg_min"], FMT_SOLES),
        ("Balde 4 kg: precio por kg con IGV - máximo", kpis["balde4_precio_kg_max"], FMT_SOLES),
        ("Margen más bajo de toda la escala (para punto de equilibrio)", kpis["margen_conservador"], FMT_PCT),
        ("   origen", kpis["margen_conservador_origen"], None),
        ("Hallazgos: alta / media / baja",
         f'{kpis["issues_alta"]} / {kpis["issues_media"]} / {kpis["issues_baja"]}', None),
        ("Hallazgos pendientes", kpis["issues_pendientes"], "0"),
    ]
    for i, (k, v, fmt) in enumerate(kpi_rows, start=3):
        ws_kpi.cell(row=i, column=1, value=k)
        c = ws_kpi.cell(row=i, column=2, value=v)
        if fmt:
            c.number_format = fmt
    ws_kpi.column_dimensions["A"].width = 58
    ws_kpi.column_dimensions["B"].width = 34

    # -- 00 Índice --
    ws_idx["A1"] = "CENTRO DE COSTOS MR. PEANUTT — todas las tablas"
    ws_idx["A1"].font = Font(bold=True, size=14, color=CACAO)
    ws_idx["A2"] = "Una hoja por tabla. Los datos salen de BD/mrpeanutt_costos.db (generada desde DATOS INICIALES/)."
    for j, h in enumerate(("Hoja", "Grupo", "Qué contiene", "Filas"), start=1):
        c = ws_idx.cell(row=4, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, (hoja, grupo, desc, n) in enumerate(indice, start=5):
        ws_idx.cell(row=i, column=1, value=hoja).hyperlink = f"#'{hoja}'!A1"
        ws_idx.cell(row=i, column=1).font = Font(color="0563C1", underline="single")
        ws_idx.cell(row=i, column=2, value=grupo)
        ws_idx.cell(row=i, column=3, value=desc)
        ws_idx.cell(row=i, column=4, value=n)
    for col, w in (("A", 28), ("B", 12), ("C", 96), ("D", 8)):
        ws_idx.column_dimensions[col].width = w

    wb.save(OUT)
    cx.close()
    print("Generado:", OUT, f"({OUT.stat().st_size} bytes, {len(indice) + 1} hojas)")


if __name__ == "__main__":
    main()
